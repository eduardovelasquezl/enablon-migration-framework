"""Construcción de `evidence_internal.xlsx` / `evidence_client.xlsx` a
partir de un `RunEvidenceContext` (ver `collector.py`).

No decide qué información mostrar por audiencia -- eso ya lo resolvió
`catalog.py` (`include_in_internal`/`include_in_client`). Este módulo solo
renderiza: pestañas, tablas, formato visual (Fase 10).
"""
from __future__ import annotations

import csv as csv_module
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .catalog import (
    CATEGORY_CATALOG,
    PROTOTYPE_LIMITATIONS_TEXT,
    SEVERITY_COLORS,
    SEVERITY_LABELS,
    CategoryDefinition,
)
from .models import RunEvidenceContext
from .sanitization import (
    category_allowed_for_audience,
    clean_numeric_display,
    relative_to_project,
    validate_audience,
)

TITLE_FONT = Font(bold=True, size=14, color="FF1F4E78")
SECTION_FONT = Font(bold=True, size=11, color="FFFFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="FF1F4E78")
LABEL_FONT = Font(bold=True, size=10)
BODY_FONT = Font(size=10)
NOTE_FONT = Font(size=10, italic=True, color="FF7F7F7F")
HEADER_FONT = Font(bold=True, size=10, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF44546A")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
WRAP_TOP_LEFT = Alignment(wrap_text=True, vertical="top", horizontal="left")
THIN_BOTTOM = Border(bottom=Side(style="thin", color="FFBFBFBF"))

DISCLAIMER_INTERNAL = (
    "Uso interno -- conserva trazabilidad técnica completa. No distribuir "
    "fuera del equipo de análisis sin revisar el contenido."
)
DISCLAIMER_CLIENT = (
    "Documento de revisión. No implica aprobación para carga en Enablon."
)

CLIENT_MOTIVO_TEMPLATES: dict[str, str] = {
    "EXCLUDED_ROWS": "Pendiente de revisión funcional -- falta información obligatoria para incluir el registro.",
    "INVALID_REFERENCE": "No existe equivalencia confirmada -- falta uno o más datos obligatorios de origen.",
    "ENTITY_UNRESOLVED": "No existe equivalencia confirmada para la entidad de origen.",
    "ENTITY_CONFLICTING": "Existe más de una equivalencia posible -- pendiente de decisión.",
    "ENTITY_EMPTY": "Información no disponible en el origen.",
    "ENTITY_DO_NOT_MIGRATE": "Clasificado como No migra según las reglas disponibles.",
    "INVALID_DATE": "Información de fecha no disponible o no reconocible en el origen.",
    "DUPLICATE_REFERENCE": "Pendiente de revisión funcional -- referencia compartida con otro registro.",
}

# Categorías cuyas incidencias por fila requieren, además de comentario,
# una decisión explícita del cliente (severidad bloqueante).
CATEGORIES_REQUIRING_CLIENT_DECISION = {"ENTITY_CONFLICTING"}

# Orden fijo de pestañas de incidencias (solo se crean las que tengan
# contenido real -- ver Fase 6/7).
ISSUE_CATEGORY_ORDER = [
    "EXCLUDED_ROWS",
    "INVALID_REFERENCE",
    "ENTITY_UNRESOLVED",
    "ENTITY_CONFLICTING",
    "ENTITY_EMPTY",
    "ENTITY_DO_NOT_MIGRATE",
    "INVALID_DATE",
    "DUPLICATE_REFERENCE",
]


# --------------------------------------------------------------------------
# Utilidades de formato (Fase 10)
# --------------------------------------------------------------------------

def _set_column_widths(ws: Worksheet, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _write_explanation_block(
    ws: Worksheet,
    category: CategoryDefinition,
    audience: str,
    record_count: int,
    start_row: int = 1,
) -> int:
    """Escribe título/descripción/severidad/impacto/responsable/acción/nº
    de registros. Devuelve la fila donde debe empezar la tabla de datos
    (deja una fila en blanco de separación, Fase 10)."""
    row = start_row
    ws.cell(row=row, column=1, value=category.title).font = TITLE_FONT
    row += 1

    description = category.description_internal if audience == "internal" else category.description_client
    cell = ws.cell(row=row, column=1, value=description)
    cell.font = BODY_FONT
    cell.alignment = WRAP_TOP_LEFT
    ws.row_dimensions[row].height = 32
    row += 1

    severity_label = SEVERITY_LABELS.get(category.severity, category.severity)
    fields = [
        ("Severidad", severity_label),
        ("Impacto en migración", category.migration_impact),
        ("Responsable", category.responsible_party),
        ("Acción requerida", category.requested_action),
        ("Número de registros", str(record_count)),
    ]
    for label, value in fields:
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        vcell = ws.cell(row=row, column=2, value=value)
        vcell.font = BODY_FONT
        vcell.alignment = WRAP_TOP_LEFT
        row += 1

    fill = PatternFill("solid", fgColor=SEVERITY_COLORS.get(category.severity, "FFFFFFFF"))
    for col in (1, 2):
        ws.cell(row=start_row + 2, column=col).fill = fill

    row += 1  # separación visual entre explicación y tabla.
    return row


def _write_data_table(
    ws: Worksheet,
    start_row: int,
    headers: list[str],
    rows: list[list],
    wrap_columns: set[int] | None = None,
) -> None:
    """Escribe cabecera + filas, con autofilter y freeze panes solo sobre
    el rango de datos (Fase 10) -- nunca sobre toda la hoja."""
    wrap_columns = wrap_columns or set()

    if not rows:
        ws.cell(row=start_row, column=1, value="(sin registros)").font = NOTE_FONT
        return

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BOTTOM
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for r_offset, row_values in enumerate(rows, start=1):
        r = start_row + r_offset
        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=r, column=col_idx, value=value)
            cell.font = BODY_FONT
            if col_idx in wrap_columns:
                cell.alignment = WRAP_TOP_LEFT

    last_row = start_row + len(rows)
    last_col_letter = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A{start_row}:{last_col_letter}{last_row}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1).coordinate

    widths = []
    for col_idx in range(1, len(headers) + 1):
        header_len = len(str(headers[col_idx - 1]))
        max_len = header_len
        for row_values in rows[:200]:  # muestreo -- suficiente para un ancho razonable sin recorrer todo en tablas grandes.
            if col_idx - 1 < len(row_values):
                max_len = max(max_len, len(str(row_values[col_idx - 1])))
        widths.append(min(max(max_len + 2, 10), 60))
    _set_column_widths(ws, widths)

    ws.sheet_view.showGridLines = False


def _new_sheet(wb: Workbook, title: str) -> Worksheet:
    ws = wb.create_sheet(title=title)
    ws.sheet_view.showGridLines = False
    return ws


# --------------------------------------------------------------------------
# Pestaña Resumen
# --------------------------------------------------------------------------

def _write_kv_rows(ws: Worksheet, row: int, pairs: list[tuple[str, str]]) -> int:
    for label, value in pairs:
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        vcell = ws.cell(row=row, column=2, value=value)
        vcell.font = BODY_FONT
        vcell.alignment = WRAP_TOP_LEFT
        row += 1
    return row


def _write_section_header(ws: Worksheet, row: int, title: str) -> int:
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = SECTION_FONT
    cell.fill = SECTION_FILL
    ws.cell(row=row, column=2).fill = SECTION_FILL
    return row + 1


def _issue_counts_by_category(ctx: RunEvidenceContext) -> dict[str, int]:
    counts: dict[str, int] = {}
    for issue in ctx.issues:
        counts[issue["category"]] = counts.get(issue["category"], 0) + 1
    return counts


def _build_summary_internal(ws: Worksheet, ctx: RunEvidenceContext) -> None:
    vr = ctx.validation_report
    run = vr.get("run", {})
    counts = vr.get("counts", {})
    ref = vr.get("reference", {})
    ent = vr.get("entities", {})
    dates = vr.get("dates", {})
    out = vr.get("output", {})
    status = vr.get("status", {})
    manifest = ctx.export_manifest

    ws.cell(row=1, column=1, value="Resumen de la ejecución -- Drills (uso interno)").font = TITLE_FONT
    row = 3

    row = _write_section_header(ws, row, "Identificación de ejecución")
    row = _write_kv_rows(ws, row, [
        ("run_id", run.get("run_id", "?")),
        ("Fecha (UTC)", run.get("timestamp", "?")),
        ("Objeto", run.get("migration_object", "?")),
        ("Modo", run.get("mode", "?")),
        ("Conexión lógica", run.get("connection_name", "?")),
        ("Versión del prototipo", manifest.get("prototype_version", "?")),
        ("Git commit", manifest.get("git_commit") or "(no disponible)"),
    ])
    row += 1

    row = _write_section_header(ws, row, "Resultado")
    row = _write_kv_rows(ws, row, [
        ("Filas leídas", str(counts.get("rows_read", 0))),
        ("Filas transformadas", str(counts.get("rows_transformed", 0))),
        ("Filas exportadas", str(counts.get("rows_exported", 0))),
        ("Filas excluidas", str(counts.get("rows_excluded", 0))),
        ("Warnings", str(counts.get("warnings", 0))),
        ("Errores", str(counts.get("errors", 0))),
        ("Estado", status.get("result", "?")),
    ])
    row += 1

    row = _write_section_header(ws, row, "Formato del CSV")
    row = _write_kv_rows(ws, row, [
        ("Encoding", out.get("encoding", "?")),
        ("BOM", str(out.get("bom", "?"))),
        ("Delimitador", repr(out.get("delimiter", "?"))),
        ("Terminador de línea", repr(out.get("line_terminator", "?"))),
        ("Número de columnas", str(out.get("column_count", "?"))),
    ])
    row += 1

    row = _write_section_header(ws, row, "Entidades (CS_ImpactedEntities)")
    row = _write_kv_rows(ws, row, [
        ("Resueltas", str(ent.get("resolved", 0))),
        ("No migra", str(ent.get("do_not_migrate", 0))),
        ("No resueltas", str(ent.get("unresolved", 0))),
        ("En conflicto", str(ent.get("conflicting", 0))),
        ("Sin información de origen", str(ent.get("empty", 0))),
    ])
    row += 1

    row = _write_section_header(ws, row, "Reference")
    row = _write_kv_rows(ws, row, [
        ("Válidas", str(ref.get("valid", 0))),
        ("Inválidas", str(ref.get("invalid", 0))),
        ("Duplicadas", str(ref.get("duplicate_references", 0))),
        ("Componente ausente -- CS_Typology", str(ref.get("missing_typology", 0))),
        ("Componente ausente -- CS_HistoricalOriginID", str(ref.get("missing_historical_origin_id", 0))),
        ("Componente ausente -- StartingDate", str(ref.get("missing_starting_date", 0))),
    ])
    row += 1

    if ctx.comparison_report and "rows" in ctx.comparison_report:
        cr = ctx.comparison_report
        row = _write_section_header(ws, row, "Comparación histórica")
        affected_cols = [c for c, v in cr.get("differences", {}).get("by_column", {}).items() if v.get("differing", 0) > 0]
        row = _write_kv_rows(ws, row, [
            ("Filas comparables (por clave)", str(cr.get("rows", {}).get("matched_by_key", 0))),
            ("Celdas coincidentes", str(cr.get("differences", {}).get("matching_cells", 0))),
            ("Celdas con diferencias", str(cr.get("differences", {}).get("differing_cells", 0))),
            ("Columnas con diferencias", ", ".join(affected_cols) if affected_cols else "(ninguna)"),
        ])
        row += 1

    row = _write_section_header(ws, row, "Estado de aprobación")
    approval_cell = ws.cell(row=row, column=1, value="approved_for_enablon_import")
    approval_cell.font = LABEL_FONT
    value_cell = ws.cell(row=row, column=2, value=str(ctx.approved_for_enablon_import).lower())
    value_cell.font = Font(bold=True, size=12, color="FF9C0006")
    value_cell.fill = PatternFill("solid", fgColor="FFF8CBAD")
    row += 2

    ws.cell(row=row, column=1, value=DISCLAIMER_INTERNAL).font = NOTE_FONT

    _set_column_widths(ws, [40, 70])


def _build_summary_client(ws: Worksheet, ctx: RunEvidenceContext) -> None:
    vr = ctx.validation_report
    run = vr.get("run", {})
    counts = vr.get("counts", {})
    status = vr.get("status", {})
    category_counts = _issue_counts_by_category(ctx)

    ws.cell(row=1, column=1, value="Resumen de la revisión -- Drills").font = TITLE_FONT
    row = 3

    row = _write_kv_rows(ws, row, [
        ("Objeto", run.get("migration_object", "?")),
        ("Fecha de la revisión", run.get("timestamp", "?")),
        ("Registros analizados", str(counts.get("rows_transformed", 0))),
        ("Registros incluidos", str(counts.get("rows_exported", 0))),
        ("Registros excluidos", str(counts.get("rows_excluded", 0))),
        ("Estado general", status.get("result", "?")),
    ])
    row += 1

    row = _write_section_header(ws, row, "Incidencias por categoría")
    if category_counts:
        for code in ISSUE_CATEGORY_ORDER:
            if code in category_counts:
                cat = CATEGORY_CATALOG[code]
                if category_allowed_for_audience(cat, "client"):
                    row = _write_kv_rows(ws, row, [(cat.title, str(category_counts[code]))])
    else:
        row = _write_kv_rows(ws, row, [("Sin incidencias detectadas", "")])
    row += 1

    row = _write_section_header(ws, row, "Acciones pendientes")
    pending = [
        CATEGORY_CATALOG[code].requested_action
        for code in ISSUE_CATEGORY_ORDER
        if code in category_counts and category_allowed_for_audience(CATEGORY_CATALOG[code], "client")
    ]
    if pending:
        for action in dict.fromkeys(pending):  # dedupe conservando orden.
            row = _write_kv_rows(ws, row, [("•", action)])
    else:
        row = _write_kv_rows(ws, row, [("•", "Ninguna acción pendiente derivada de incidencias de datos.")])
    row += 1

    ws.cell(row=row, column=1, value=DISCLAIMER_CLIENT).font = Font(bold=True, italic=True, size=10, color="FF9C0006")

    _set_column_widths(ws, [45, 65])


# --------------------------------------------------------------------------
# Pestaña Guía
# --------------------------------------------------------------------------

def _build_guide(ws: Worksheet, audience: str) -> None:
    ws.cell(row=1, column=1, value="Guía de lectura").font = TITLE_FONT
    row = 3

    row = _write_section_header(ws, row, "Qué es este documento")
    row = _write_kv_rows(ws, row, [
        ("", "Evidencia generada automáticamente tras una ejecución del prototipo de "
             "exportación de Drills. Resume el resultado, las incidencias detectadas y "
             "las preguntas todavía abiertas."),
    ])
    row += 1

    row = _write_section_header(ws, row, "Severidades")
    for key, label in SEVERITY_LABELS.items():
        meaning = {
            "info": "Informativo -- no requiere ninguna acción.",
            "review_required": "Revisión requerida -- no es un error de carga confirmado, pero necesita confirmación funcional.",
            "blocking_for_approval": "Bloqueante -- impide aprobar la especificación hasta resolverse.",
            "not_migrated_by_design": "No migra -- exclusión esperada y ya acordada, no es un fallo técnico.",
        }.get(key, "")
        row = _write_kv_rows(ws, row, [(label, meaning)])
    row += 1

    row = _write_section_header(ws, row, "Cómo interpretar 'No migra'")
    row = _write_kv_rows(ws, row, [
        ("", "Un registro clasificado 'No migra' no es un error técnico. Corresponde a "
             "una regla ya acordada que excluye determinadas entidades de la carga -- "
             "se documenta para trazabilidad, no para corrección."),
    ])
    row += 1

    row = _write_section_header(ws, row, "Error, warning y revisión funcional")
    row = _write_kv_rows(ws, row, [
        ("Error", "Impide completar la ejecución o escribir el resultado."),
        ("Warning", "La ejecución se completa, pero hay un dato que no se ha podido resolver del todo."),
        ("Revisión funcional", "El dato se ha procesado sin fallo técnico, pero requiere una confirmación de negocio antes de aprobarse."),
    ])
    row += 1

    row = _write_section_header(ws, row, "Sobre el CSV histórico")
    row = _write_kv_rows(ws, row, [
        ("", "El CSV histórico usado como referencia de comparación es una exportación "
             "de datos reales, no una plantilla oficial de importación de Enablon."),
    ])
    row += 1

    if audience == "internal":
        row = _write_section_header(ws, row, "Significado de cada pestaña")
        row = _write_kv_rows(ws, row, [
            (CATEGORY_CATALOG[code].short_title, CATEGORY_CATALOG[code].title)
            for code in ["EXPORT_SUMMARY", *ISSUE_CATEGORY_ORDER, "HISTORICAL_DIFFERENCES"]
        ])

    _set_column_widths(ws, [30, 80])


# --------------------------------------------------------------------------
# Pestaña Hallazgos
# --------------------------------------------------------------------------

def _build_findings(ws: Worksheet, ctx: RunEvidenceContext, audience: str) -> None:
    ws.cell(row=1, column=1, value="Hallazgos funcionales").font = TITLE_FONT
    row = 3

    for code in ("STARTING_DATE_TIME_GAP", "ENTITY_CATALOG_FIDELITY"):
        cat = CATEGORY_CATALOG[code]
        row = _write_section_header(ws, row, cat.title)
        description = cat.description_internal if audience == "internal" else cat.description_client
        row = _write_kv_rows(ws, row, [
            ("Descripción", description),
            ("Severidad", SEVERITY_LABELS.get(cat.severity, cat.severity)),
            ("Impacto en migración", cat.migration_impact),
            ("Responsable propuesto", cat.responsible_party),
            ("Acción requerida", cat.requested_action),
        ])
        row += 1

    row = _write_section_header(ws, row, "Preguntas abiertas")
    for qid, data in ctx.open_questions.items():
        if data is None:
            row = _write_kv_rows(ws, row, [
                (qid, "Referencia documental: docs/specifications/v1.0/export/open_questions.md (fila no localizada en este incremento)."),
            ])
            continue
        summary = (
            f"Impacto: {data['impact']} | Responsable: {data['responsible']} | Estado: {data['status']}"
            if audience == "internal"
            else f"Responsable: {data['responsible']} | Estado: {data['status']}"
        )
        row = _write_kv_rows(ws, row, [(qid, summary)])
    row += 1

    row = _write_section_header(ws, row, "Limitaciones del prototipo")
    for text in PROTOTYPE_LIMITATIONS_TEXT:
        row = _write_kv_rows(ws, row, [("•", text)])

    _set_column_widths(ws, [22, 95])


# --------------------------------------------------------------------------
# Pestañas de detalle (a partir de issues.jsonl)
# --------------------------------------------------------------------------

def _internal_issue_table(category_code: str, rows: list[dict]) -> tuple[list[str], list[list]]:
    headers = ["ID histórico", "Origen", "Valor fuente", "Valor mapeado", "Código", "Evidencia", "Detalle técnico"]
    table = []
    for r in rows:
        table.append([
            r.get("historical_origin_id") or "",
            r.get("historical_data_origin") or "",
            clean_numeric_display(r.get("source_value")),
            clean_numeric_display(r.get("mapped_value")),
            r.get("category", category_code),
            r.get("evidence_id") or "",
            r.get("message") or "",
        ])
    return headers, table


def _client_issue_table(category_code: str, rows: list[dict]) -> tuple[list[str], list[list]]:
    needs_decision = category_code in CATEGORIES_REQUIRING_CLIENT_DECISION
    headers = ["Identificador histórico", "Valor actual", "Resultado de mapping", "Motivo", "Acción solicitada", "Comentario cliente"]
    if needs_decision:
        headers.append("Decisión cliente")
    action = CATEGORY_CATALOG[category_code].requested_action
    motivo = CLIENT_MOTIVO_TEMPLATES.get(category_code, "Requiere revisión funcional.")

    table = []
    for r in rows:
        row_values = [
            r.get("historical_origin_id") or "",
            clean_numeric_display(r.get("source_value")),
            clean_numeric_display(r.get("mapped_value")) or "(sin resultado)",
            motivo,
            action,
            "",  # Comentario cliente -- vacío a propósito.
        ]
        if needs_decision:
            row_values.append("")  # Decisión cliente -- vacío a propósito.
        table.append(row_values)
    return headers, table


def _build_issue_sheet(wb: Workbook, category_code: str, rows: list[dict], audience: str) -> None:
    category = CATEGORY_CATALOG[category_code]
    ws = _new_sheet(wb, category.short_title)
    next_row = _write_explanation_block(ws, category, audience, len(rows))

    if audience == "internal":
        headers, table = _internal_issue_table(category_code, rows)
        wrap_cols = {7}
    else:
        headers, table = _client_issue_table(category_code, rows)
        wrap_cols = {4}

    _write_data_table(ws, next_row, headers, table, wrap_columns=wrap_cols)


def _build_historical_differences_sheet(wb: Workbook, ctx: RunEvidenceContext, audience: str) -> None:
    category = CATEGORY_CATALOG["HISTORICAL_DIFFERENCES"]
    by_column = (ctx.comparison_report or {}).get("differences", {}).get("by_column", {})
    ws = _new_sheet(wb, category.short_title)
    next_row = _write_explanation_block(ws, category, audience, len(by_column))

    headers = ["Columna", "Coincidencias", "Diferencias", "% coincidencia"]
    table = []
    for col, counts in sorted(by_column.items()):
        matching = counts.get("matching", 0)
        differing = counts.get("differing", 0)
        total = matching + differing
        pct = f"{(matching / total * 100):.0f}%" if total else "N/A"
        table.append([col, matching, differing, pct])

    _write_data_table(ws, next_row, headers, table)


def _build_exported_sheet(wb: Workbook, ctx: RunEvidenceContext) -> None:
    """Solo versión interna -- lee `drills.csv` directamente (no duplica su
    contenido en `issues.jsonl`, ver Fase 8)."""
    category = CATEGORY_CATALOG["EXPORTED_ROWS"]
    ws = _new_sheet(wb, category.short_title)

    with open(ctx.csv_path, encoding="utf-8", newline="") as f:
        reader = csv_module.reader(f, delimiter="\t")
        csv_rows = list(reader)
    headers = csv_rows[0] if csv_rows else []
    data = csv_rows[1:]

    next_row = _write_explanation_block(ws, category, "internal", len(data))
    _write_data_table(ws, next_row, headers, data)


# --------------------------------------------------------------------------
# Pestaña Trazabilidad (solo interna)
# --------------------------------------------------------------------------

def _build_traceability(ws: Worksheet, ctx: RunEvidenceContext) -> None:
    manifest = ctx.export_manifest
    ws.cell(row=1, column=1, value="Trazabilidad técnica").font = TITLE_FONT
    row = 3

    row = _write_section_header(ws, row, "Identificadores de evidencia")
    for evidence_id in manifest.get("evidence_ids", []):
        row = _write_kv_rows(ws, row, [("evidence_id", evidence_id)])
    row += 1

    hashes = manifest.get("hashes", {})
    row = _write_section_header(ws, row, "Hashes SHA-256")
    row = _write_kv_rows(ws, row, [
        ("SQL", hashes.get("sql_sha256", "?")),
        ("Mappings", hashes.get("mappings_sha256", "?")),
        ("Configuración", hashes.get("config_sha256", "?")),
        ("CSV generado", hashes.get("csv_sha256", "?")),
    ])
    row += 1

    row = _write_section_header(ws, row, "Rutas relativas de artefactos")
    row = _write_kv_rows(ws, row, [
        ("SQL de origen", manifest.get("source", {}).get("sql_file", "?")),
        ("drills.csv", relative_to_project(ctx.csv_path)),
        ("validation_report.yaml", relative_to_project(ctx.run_dir / "validation_report.yaml")),
        ("export_manifest.yaml", relative_to_project(ctx.run_dir / "export_manifest.yaml")),
    ])
    row += 1

    row = _write_section_header(ws, row, "Reglas aplicadas")
    for rule in manifest.get("rules_applied", []):
        row = _write_kv_rows(ws, row, [("•", rule)])
    row += 1

    ws.cell(
        row=row, column=1,
        value="No se incluyen contraseñas ni cadenas de conexión completas -- "
              "ver export_manifest.yaml -> connection.note.",
    ).font = NOTE_FONT

    _set_column_widths(ws, [28, 90])


# --------------------------------------------------------------------------
# Orquestación
# --------------------------------------------------------------------------

def build_workbook(ctx: RunEvidenceContext, audience: str) -> Workbook:
    validate_audience(audience)

    wb = Workbook()
    wb.remove(wb.active)  # la hoja por defecto de openpyxl no se usa.

    summary_ws = _new_sheet(wb, "Resumen")
    if audience == "internal":
        _build_summary_internal(summary_ws, ctx)
    else:
        _build_summary_client(summary_ws, ctx)

    guide_ws = _new_sheet(wb, "Guía")
    _build_guide(guide_ws, audience)

    findings_ws = _new_sheet(wb, "Hallazgos")
    _build_findings(findings_ws, ctx, audience)

    issues_by_category: dict[str, list[dict]] = {}
    for issue in ctx.issues:
        issues_by_category.setdefault(issue["category"], []).append(issue)

    for code in ISSUE_CATEGORY_ORDER:
        rows = issues_by_category.get(code, [])
        category = CATEGORY_CATALOG[code]
        if not rows:
            continue  # nunca se crean pestañas vacías (Fase 10/11).
        if not category_allowed_for_audience(category, audience):
            continue
        _build_issue_sheet(wb, code, rows, audience)

    if ctx.comparison_report and (ctx.comparison_report.get("differences", {}) or {}).get("by_column"):
        if category_allowed_for_audience(CATEGORY_CATALOG["HISTORICAL_DIFFERENCES"], audience):
            _build_historical_differences_sheet(wb, ctx, audience)

    if audience == "internal":
        _build_exported_sheet(wb, ctx)
        traceability_ws = _new_sheet(wb, "Trazabilidad")
        _build_traceability(traceability_ws, ctx)

    return wb


def save_workbook(wb: Workbook, path: Path) -> Path:
    """Escritura atómica (temporal + `os.replace`), mismo patrón que el
    resto de artefactos del pipeline."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=str(path.parent))
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, path)
    except BaseException:
        if tmp_path.exists():
            tmp_path.unlink()
        raise
    return path
