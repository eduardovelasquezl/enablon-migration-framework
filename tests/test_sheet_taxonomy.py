"""Tests de la taxonomía ampliada de hojas ETL (src/analysis/_sheet_taxonomy.py)
y de su exposición pública en schema_analyzer.py. Ejecutar con: pytest tests/

Todos los Excel de prueba se generan dinámicamente en tmp_path."""
import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.analysis._sheet_taxonomy import classify_sheet
from src.analysis.schema_analyzer import classify_workbook_sections

FIELD_MAPPING_HEADER = ["CampoOrigen", "Field Destiny XML", "Field Destiny ES", "Adaptación", "Transformation From", None, "XML", "ES"]
RULE_HEADER = ["DatoOrigen", "DatoDestino", "EsCondicion", "ReglaEspecial", "Parametro"]


# ---------------------------------------------------------------------------
# Hoja de mapping -- confirmed / alta confianza
# ---------------------------------------------------------------------------

def test_hoja_de_field_mapping_confirmed_alta_confianza():
    classifications = classify_sheet(
        sheet_title="MapeoBypass",
        headers=FIELD_MAPPING_HEADER,
        max_row=76,
        source_file="etl.xlsx",
    )
    principal = classifications[0]
    assert principal.category == "field_mapping"
    assert principal.evidence_status == "confirmed"
    assert principal.confidence_score == 100


# ---------------------------------------------------------------------------
# Catálogo de referencia Enablon -- segunda sección de la MISMA hoja
# ---------------------------------------------------------------------------

def test_hoja_mixta_produce_dos_clasificaciones_mapping_y_catalogo():
    classifications = classify_sheet(
        sheet_title="MapeoBypass",
        headers=FIELD_MAPPING_HEADER,
        max_row=76,
        source_file="etl.xlsx",
    )
    categories = {c.category for c in classifications}
    assert categories == {"field_mapping", "enablon_reference_catalog"}

    catalog = next(c for c in classifications if c.category == "enablon_reference_catalog")
    assert catalog.evidence_status == "confirmed"
    assert "G-H" in catalog.approximate_range


def test_hoja_sin_columnas_xml_es_no_produce_catalogo():
    classifications = classify_sheet(
        sheet_title="MapeoSinCatalogo",
        headers=FIELD_MAPPING_HEADER[:6],  # sin columnas G/H
        max_row=10,
        source_file="etl.xlsx",
    )
    assert len(classifications) == 1
    assert classifications[0].category == "field_mapping"


# ---------------------------------------------------------------------------
# Equivalencias (DatoOrigen/DatoDestino)
# ---------------------------------------------------------------------------

def test_hoja_de_equivalencias():
    classifications = classify_sheet(
        sheet_title="Mapeo_Nivel",
        headers=RULE_HEADER,
        max_row=21,
        source_file="etl.xlsx",
    )
    assert classifications[0].category == "equivalence_table"
    assert classifications[0].evidence_status == "inferred"


def test_hoja_con_nombre_de_regla_conocida_es_transformation_rule():
    classifications = classify_sheet(
        sheet_title="CharacterFix",
        headers=[None, None],  # sin cabecera reconocible -- se clasifica por nombre
        max_row=2,
        source_file="etl.xlsx",
    )
    assert classifications[0].category == "transformation_rule"
    assert classifications[0].evidence_status == "inferred"
    assert classifications[0].confidence_score == 90


# ---------------------------------------------------------------------------
# CSV preview
# ---------------------------------------------------------------------------

def test_hoja_csv_preview_por_nombre():
    classifications = classify_sheet(
        sheet_title="CSV_SM_FULL",
        headers=["CS_Entity", "StartDate"],
        max_row=10253,
        source_file="etl.xlsx",
    )
    # Con >200 filas y nombre CSV_*, el nombre gana porque se evalúa antes
    # que el volumen -- confirma que no se cae en "source_data" por defecto.
    assert classifications[0].category == "csv_preview"
    assert classifications[0].evidence_status == "inferred"
    assert classifications[0].confidence_score == 40


# ---------------------------------------------------------------------------
# Hoja realmente desconocida -- pending_classification, nunca forzada
# ---------------------------------------------------------------------------

def test_hoja_desconocida_queda_pending_no_se_fuerza():
    classifications = classify_sheet(
        sheet_title="Hoja_XYZ123_Temporal",
        headers=["ColumnaA", "ColumnaB"],
        max_row=5,
        source_file="etl.xlsx",
    )
    assert classifications[0].category == "pending_classification"
    assert classifications[0].evidence_status == "pending"
    assert classifications[0].confidence_score == 0
    assert classifications[0].warnings


def test_hoja_grande_sin_senal_es_source_data():
    classifications = classify_sheet(
        sheet_title="ITP_SIMULACRO",
        headers=["IDSIM", "IDCentro"],
        max_row=12501,
        source_file="etl.xlsx",
    )
    assert classifications[0].category == "source_data"
    assert classifications[0].evidence_status == "inferred"


# ---------------------------------------------------------------------------
# confidence_score y evidence_status son ejes independientes
# ---------------------------------------------------------------------------

def test_confidence_alto_no_implica_evidence_status_confirmed():
    # Nombre de regla conocido: confianza alta (90) pero NUNCA "confirmed"
    # solo por el nombre -- se necesitaría inspeccionar contenido para eso.
    classifications = classify_sheet(
        sheet_title="titlefix",
        headers=[None],
        max_row=2,
        source_file="etl.xlsx",
    )
    c = classifications[0]
    assert c.confidence_score == 90
    assert c.evidence_status == "inferred"  # no "confirmed"


def test_categoria_invalida_lanza_error():
    from src.analysis._sheet_taxonomy import SheetClassification
    import pytest
    with pytest.raises(ValueError):
        SheetClassification(category="no_existe", evidence_status="confirmed", confidence_score=100)


def test_confidence_score_fuera_de_rango_lanza_error():
    from src.analysis._sheet_taxonomy import SheetClassification
    import pytest
    with pytest.raises(ValueError):
        SheetClassification(category="index", evidence_status="confirmed", confidence_score=150)


# ---------------------------------------------------------------------------
# Marcador "No migra" en contenido
# ---------------------------------------------------------------------------

def test_marcador_no_migra_detectado_por_contenido():
    classifications = classify_sheet(
        sheet_title="Entidades_Excluidas",
        headers=["Code", "Motivo"],
        max_row=10,
        source_file="etl.xlsx",
        contains_no_migrate_marker=True,
    )
    assert classifications[0].category == "exclusions_no_migrate"
    assert classifications[0].evidence_status == "inferred"
    assert classifications[0].warnings


# ---------------------------------------------------------------------------
# Superficie pública: schema_analyzer.classify_workbook_sections
# ---------------------------------------------------------------------------

def test_classify_workbook_sections_expone_taxonomia_ampliada(tmp_path):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Index"
    ws1.append(["SOURCES DATA", "MAPS", "FINAL DATA"])

    ws2 = wb.create_sheet("MapeoBypass")
    ws2.append(FIELD_MAPPING_HEADER)
    ws2.append(["IDBES", None, "Historical Origin ID", "Si", "map_BES_Fix", None, "CS_HistoricalOriginID", "Historical Origin ID"])

    ws3 = wb.create_sheet("HojaMisteriosa")
    ws3.append(["ColX", "ColY"])
    ws3.append([1, 2])

    path = tmp_path / "etl.xlsx"
    wb.save(path)

    result = classify_workbook_sections(path)
    assert result["Index"][0].category == "index"
    categories_mapeo = {c.category for c in result["MapeoBypass"]}
    assert categories_mapeo == {"field_mapping", "enablon_reference_catalog"}
    assert result["HojaMisteriosa"][0].category == "pending_classification"


def test_classify_workbook_sections_con_no_migrate_marker(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Exclusiones"
    ws.append(["Code", "Estado"])
    ws.append(["ENERGYP.SR", "No migra"])
    path = tmp_path / "etl.xlsx"
    wb.save(path)

    result_off = classify_workbook_sections(path, check_no_migrate_marker=False)
    assert result_off["Exclusiones"][0].category != "exclusions_no_migrate"

    result_on = classify_workbook_sections(path, check_no_migrate_marker=True)
    assert result_on["Exclusiones"][0].category == "exclusions_no_migrate"
