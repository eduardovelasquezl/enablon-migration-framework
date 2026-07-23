"""
Lee las hojas de mapeo (patrón `CampoOrigen | Field Destiny XML | ... |
Adaptación | Transformation From | ...` y el patrón de regla `DatoOrigen |
DatoDestino | EsCondicion | ReglaEspecial | Parametro`) de un ETL original y
las convierte en estructuras declarativas simples, para que
mapping_engine.py no tenga que tocar Excel directamente.
"""
from __future__ import annotations

from dataclasses import dataclass

from openpyxl import load_workbook


@dataclass
class FieldMapping:
    campo_origen: str
    campo_destino_es: str | None
    requiere_adaptacion: bool
    regla: str | None
    campo_destino_xml: str | None


@dataclass
class RuleTableEntry:
    dato_origen: str
    dato_destino: str
    es_condicion: str | None
    regla_especial: str | None
    parametro: str | None


def read_field_mapping_sheet(path: str, sheet_name: str) -> list[FieldMapping]:
    """Lee una hoja del patrón de mapeo de campo (8 columnas conocidas).

    Nota de diseño confirmada en el proyecto original: la columna B
    (`Field Destiny XML`) suele ser una fórmula XLOOKUP que solo VALIDA que
    el campo destino existe en el catálogo de Enablon — la transformación
    real vive en la columna `Transformation From` (nombre de la
    hoja de regla a aplicar). No asumir que B ya trae el valor transformado.
    """
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        return []

    mappings = []
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        campo_origen = row[0]
        campo_destino_es = row[2] if len(row) > 2 else None
        adaptacion_raw = str(row[3]).strip().lower() if len(row) > 3 and row[3] is not None else "no"
        regla = row[4] if len(row) > 4 else None
        campo_destino_xml = row[6] if len(row) > 6 else None
        mappings.append(
            FieldMapping(
                campo_origen=str(campo_origen),
                campo_destino_es=str(campo_destino_es) if campo_destino_es else None,
                requiere_adaptacion=adaptacion_raw in ("sí", "si", "yes", "true"),
                regla=str(regla) if regla else None,
                campo_destino_xml=str(campo_destino_xml) if campo_destino_xml else None,
            )
        )
    return mappings


def read_rule_table_sheet(path: str, sheet_name: str) -> list[RuleTableEntry]:
    """Lee una hoja de regla pequeña (patrón DatoOrigen/DatoDestino/...)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        return []

    entries = []
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        entries.append(
            RuleTableEntry(
                dato_origen=str(row[0]),
                dato_destino=str(row[1]) if len(row) > 1 else "",
                es_condicion=str(row[2]) if len(row) > 2 and row[2] is not None else None,
                regla_especial=str(row[3]) if len(row) > 3 and row[3] is not None else None,
                parametro=str(row[4]) if len(row) > 4 and row[4] is not None else None,
            )
        )
    return entries


def rule_table_to_dict(entries: list[RuleTableEntry]) -> dict[str, str]:
    """Convierte una hoja de regla en un diccionario simple DatoOrigen ->
    DatoDestino, para usar directamente con `lookup_simple`.

    Trata el valor comodín '*' (visto en varios módulos como regla ELSE)
    como clave especial "__default__".
    """
    out = {}
    for e in entries:
        key = "__default__" if e.dato_origen.strip() == "*" else e.dato_origen
        out[key] = e.dato_destino
    return out
