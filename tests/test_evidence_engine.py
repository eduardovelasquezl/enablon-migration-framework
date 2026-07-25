"""Tests del Evidence Engine v0.1 (sin red, sin SQL Server real -- usan una
ejecución sintética con artefactos ya escritos, exactamente como los
escribiría `src.export.prototype.drills.pipeline`).

Ejecutar con: pytest tests/test_evidence_engine.py -v
"""
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
import yaml
from click.testing import CliRunner
from openpyxl import load_workbook

from src.cli import cli
from src.evidence.catalog import CATEGORY_CATALOG
from src.evidence.collector import load_run, resolve_run_dir
from src.evidence.models import EvidenceSourceError
from src.evidence.workbook import build_workbook, save_workbook

CSV_COLUMNS = [
    "CS_Typology", "Reference", "StartingDate", "CS_HistoricalOriginID",
    "CS_Letter", "CS_ImpactedEntities", "CS_WorkflowStatus", "CS_HistoricalDataOrigin",
]

FIXTURE_ISSUES = [
    {"run_id": "fx000000aaaa", "row_key": "101", "historical_origin_id": "101",
     "historical_data_origin": "Prevención.ITP_SIM+ITP_SIM_ACCIONES_CORRECTORAS",
     "category": "EXCLUDED_ROWS", "severity": "review_required", "source_value": None,
     "mapped_value": None, "message": "Fila excluida -- missing_starting_date.",
     "evidence_id": "AFD-DRILLS-REFERENCE-001", "included_in_csv": False},
    {"run_id": "fx000000aaaa", "row_key": "101", "historical_origin_id": "101",
     "historical_data_origin": "Prevención.ITP_SIM+ITP_SIM_ACCIONES_CORRECTORAS",
     "category": "INVALID_REFERENCE", "severity": "review_required", "source_value": "missing_starting_date",
     "mapped_value": None, "message": "No fue posible construir Reference.",
     "evidence_id": "AFD-DRILLS-REFERENCE-001", "included_in_csv": False},
    {"run_id": "fx000000aaaa", "row_key": "101", "historical_origin_id": "101",
     "historical_data_origin": "Prevención.ITP_SIM+ITP_SIM_ACCIONES_CORRECTORAS",
     "category": "INVALID_DATE", "severity": "review_required", "source_value": None,
     "mapped_value": None, "message": "Fecha ausente en el origen.",
     "evidence_id": "evidence:sql_source.simulacros_dataset", "included_in_csv": False},
    {"run_id": "fx000000aaaa", "row_key": "202", "historical_origin_id": "202",
     "historical_data_origin": "Prevención.ITP_SIM+ITP_SIM_ACCIONES_CORRECTORAS",
     "category": "ENTITY_UNRESOLVED", "severity": "review_required", "source_value": "999999.0",
     "mapped_value": None, "message": "IDUnidadOrg no está en el catálogo de entidad normalizado.",
     "evidence_id": "object_assessments/drills_entity_resolution_assessment.md", "included_in_csv": True},
    {"run_id": "fx000000aaaa", "row_key": "303", "historical_origin_id": "303",
     "historical_data_origin": "Prevención.ITP_SIM+ITP_SIM_ACCIONES_CORRECTORAS",
     "category": "ENTITY_CONFLICTING", "severity": "blocking_for_approval", "source_value": "278.0",
     "mapped_value": None, "message": "IDUnidadOrg tiene más de un Code distinto en el catálogo.",
     "evidence_id": "object_assessments/drills_entity_resolution_assessment.md", "included_in_csv": True},
    {"run_id": "fx000000aaaa", "row_key": "404", "historical_origin_id": "404",
     "historical_data_origin": "Prevención.ITP_SIM+ITP_SIM_ACCIONES_CORRECTORAS",
     "category": "ENTITY_EMPTY", "severity": "review_required", "source_value": None,
     "mapped_value": None, "message": "IDUnidadOrg ausente en el origen.",
     "evidence_id": "object_assessments/drills_entity_resolution_assessment.md", "included_in_csv": True},
    {"run_id": "fx000000aaaa", "row_key": "505", "historical_origin_id": "505",
     "historical_data_origin": "Prevención.ITP_SIM+ITP_SIM_ACCIONES_CORRECTORAS",
     "category": "ENTITY_DO_NOT_MIGRATE", "severity": "not_migrated_by_design", "source_value": "521.0",
     "mapped_value": "No migra", "message": "IDUnidadOrg resuelve a una entidad 'No migra'.",
     "evidence_id": "object_assessments/drills_entity_resolution_assessment.md", "included_in_csv": True},
    {"run_id": "fx000000aaaa", "row_key": "606", "historical_origin_id": "606",
     "historical_data_origin": "Prevención.ITP_SIM+ITP_SIM_ACCIONES_CORRECTORAS",
     "category": "DUPLICATE_REFERENCE", "severity": "review_required", "source_value": None,
     "mapped_value": "PEI-HIST-606-01/01/2020", "message": "Reference duplicada.",
     "evidence_id": "AFD-DRILLS-REFERENCE-001", "included_in_csv": True},
    {"run_id": "fx000000aaaa", "row_key": "607", "historical_origin_id": "607",
     "historical_data_origin": "Prevención.ITP_SIM+ITP_SIM_ACCIONES_CORRECTORAS",
     "category": "DUPLICATE_REFERENCE", "severity": "review_required", "source_value": None,
     "mapped_value": "PEI-HIST-606-01/01/2020", "message": "Reference duplicada.",
     "evidence_id": "AFD-DRILLS-REFERENCE-001", "included_in_csv": True},
]


def _build_fixture_run(base_dir: Path, run_id: str = "fx000000aaaa") -> Path:
    """Construye una carpeta de ejecución sintética con la misma forma
    exacta que escribe `src.export.prototype.drills.pipeline.run()`, pero
    sin tocar SQL Server -- para poder probar TODAS las categorías de
    incidencia en un solo run, algo que una ejecución real no garantiza."""
    run_dir = base_dir / "20260101T000000Z"
    run_dir.mkdir(parents=True)

    validation_report = {
        "run": {
            "run_id": run_id, "timestamp": "20260101T000000Z", "mode": "sample",
            "connection_name": "prevencion", "migration_object": "Drills",
            "prototype_status": "review_only",
        },
        "counts": {
            "rows_read": 10, "rows_transformed": 10, "rows_exported": 7,
            "rows_excluded": 3, "warnings": 5, "errors": 0,
        },
        "reference": {
            "valid": 7, "invalid": 3, "missing_typology": 0,
            "missing_historical_origin_id": 0, "missing_starting_date": 3,
            "duplicate_references": 2,
        },
        "entities": {"resolved": 3, "do_not_migrate": 1, "unresolved": 1, "conflicting": 1, "empty": 1},
        "dates": {"valid": 7, "invalid": 0, "empty": 3},
        "output": {
            "path": str(run_dir / "drills.csv"), "encoding": "utf-8", "bom": False,
            "delimiter": "\t", "line_terminator": "\\r\\n", "column_count": 8,
            "columns": CSV_COLUMNS,
        },
        "status": {"result": "SUCCESS_WITH_WARNINGS", "blocking_errors": [], "warnings": ["5 warnings de ejemplo"]},
    }
    (run_dir / "validation_report.yaml").write_text(
        yaml.safe_dump(validation_report, allow_unicode=True, sort_keys=False), encoding="utf-8"
    )

    export_manifest = {
        "run_id": run_id, "timestamp": "20260101T000000Z", "git_commit": None,
        "migration_object": "Drills", "module": "simulacros", "prototype_version": "0.1.0-prototype",
        "prototype_status": "review_only", "mode": "sample",
        "connection": {"name": "prevencion", "note": "Sin credenciales."},
        "source": {
            "sql_file": "sql/source_queries/Simulacros/SQLQuery - DATASET SIMULACRO.sql",
            "sql_sha256": "a" * 64, "sql_modified_from_original": False,
            "evidence_id": "evidence:sql_source.simulacros_dataset",
        },
        "hashes": {"sql_sha256": "a" * 64, "mappings_sha256": "b" * 64, "config_sha256": "c" * 64, "csv_sha256": "d" * 64},
        "output": {"row_count": 7, "columns": CSV_COLUMNS, "encoding": "utf-8", "delimiter": "\t", "bom": False, "line_terminator": "\r\n"},
        "rules_applied": ["build_reference (AFD-DRILLS-REFERENCE-001)"],
        "evidence_ids": ["AFD-DRILLS-REFERENCE-001", "evidence:sql_source.simulacros_dataset"],
        "limitations": ["Prototipo limitado a 8 de 36 columnas."],
        "open_questions": ["OQ-ETL-05", "OQ-ETL-06", "OQ-ENT-04", "OQ-ETL-02", "OQ-OBJ-07"],
        "approved_for_enablon_import": False,
    }
    (run_dir / "export_manifest.yaml").write_text(
        yaml.safe_dump(export_manifest, allow_unicode=True, sort_keys=False), encoding="utf-8"
    )

    comparison_report = {
        "baseline": {"path": "inputs/.../Drills-22072026-41.csv", "role": "historical_output_evidence", "not_an_official_template": True},
        "schema": {"matching_columns": CSV_COLUMNS, "missing_in_generated": ["NameEN"], "additional_in_generated": [], "order_matches": False},
        "rows": {"historical": 11921, "generated": 7, "matched_by_key": 5, "only_historical": 11916, "only_generated": 2},
        "differences": {
            "total_cells_compared": 40, "matching_cells": 32, "differing_cells": 8,
            "by_column": {
                "CS_ImpactedEntities": {"matching": 2, "differing": 3},
                "StartingDate": {"matching": 1, "differing": 4},
                "Reference": {"matching": 5, "differing": 0},
            },
        },
        "reference": {"matching": 5, "differing": 0, "historical_invalid": 0, "generated_invalid": 0},
        "limitations": ["Comparación de ejemplo."],
    }
    (run_dir / "comparison_report.yaml").write_text(
        yaml.safe_dump(comparison_report, allow_unicode=True, sort_keys=False), encoding="utf-8"
    )

    with open(run_dir / "issues.jsonl", "w", encoding="utf-8") as f:
        for issue in FIXTURE_ISSUES:
            f.write(json.dumps(issue, ensure_ascii=False) + "\n")

    csv_rows = [CSV_COLUMNS]
    for i in range(7):
        csv_rows.append([f"PEI", f"PEI-HIST-{600+i}-01/01/2020", "01/01/2020 00:00", str(600 + i), "A", "MCPF.HIS", "Validated", "Prevención.ITP_SIM+ITP_SIM_ACCIONES_CORRECTORAS"])
    csv_text = "\n".join("\t".join(row) for row in csv_rows) + "\n"
    (run_dir / "drills.csv").write_text(csv_text, encoding="utf-8", newline="")

    return run_dir


@pytest.fixture()
def fixture_run(tmp_path):
    return _build_fixture_run(tmp_path)


@pytest.fixture()
def internal_wb(fixture_run):
    ctx = load_run(fixture_run)
    return build_workbook(ctx, "internal")


@pytest.fixture()
def client_wb(fixture_run):
    ctx = load_run(fixture_run)
    return build_workbook(ctx, "client")


# --------------------------------------------------------------------------
# 1-2. Generación de ambos Excel
# --------------------------------------------------------------------------

def test_genera_excel_interno(fixture_run, tmp_path):
    ctx = load_run(fixture_run)
    wb = build_workbook(ctx, "internal")
    path = save_workbook(wb, tmp_path / "evidence_internal.xlsx")
    assert path.is_file()


def test_genera_excel_cliente(fixture_run, tmp_path):
    ctx = load_run(fixture_run)
    wb = build_workbook(ctx, "client")
    path = save_workbook(wb, tmp_path / "evidence_client.xlsx")
    assert path.is_file()


# --------------------------------------------------------------------------
# 3. Regeneración desde un run existente sin SQL
# --------------------------------------------------------------------------

def test_regeneracion_no_toca_sql(fixture_run, monkeypatch):
    """Si `load_run`/`build_workbook` intentaran ejecutar SQL, esto
    fallaría (no hay conexión real disponible en un test unitario) -- el
    hecho de que pase confirma que no se ejecuta ninguna query."""
    import src.db.query_runner as qr

    def _forbidden(*args, **kwargs):
        raise AssertionError("No debería ejecutarse ninguna query SQL al regenerar evidencia.")

    monkeypatch.setattr(qr, "run_query", _forbidden)
    ctx = load_run(fixture_run)
    build_workbook(ctx, "both") if False else None  # 'both' se resuelve en la CLI, no aquí.
    build_workbook(ctx, "internal")
    build_workbook(ctx, "client")


# --------------------------------------------------------------------------
# 4-6. Existencia de Resumen / Guía / Hallazgos
# --------------------------------------------------------------------------

def test_pestanas_obligatorias_presentes(internal_wb, client_wb):
    for wb in (internal_wb, client_wb):
        for required in ("Resumen", "Guía", "Hallazgos"):
            assert required in wb.sheetnames


# --------------------------------------------------------------------------
# 7 / 20. Ausencia de pestañas vacías
# --------------------------------------------------------------------------

def test_no_hay_pestanas_de_incidencia_sin_registros(internal_wb):
    # El fixture no genera ninguna incidencia INVALID_DATE con contenido
    # separado de EXCLUDED_ROWS salvo las ya declaradas -- toda pestaña de
    # categoría presente debe tener al menos una fila de datos.
    for sheet_name in internal_wb.sheetnames:
        if sheet_name in ("Resumen", "Guía", "Hallazgos", "Trazabilidad", "Exportados"):
            continue
        ws = internal_wb[sheet_name]
        # fila 1 = título, deben existir filas de tabla más abajo (no solo "(sin registros)").
        values = [c.value for row in ws.iter_rows() for c in row if c.value]
        assert not any(v == "(sin registros)" for v in values), f"{sheet_name} no debería existir sin registros"


# --------------------------------------------------------------------------
# 8. Nombres de pestaña válidos y <= 31 caracteres
# --------------------------------------------------------------------------

def test_nombres_de_pestana_validos(internal_wb, client_wb):
    invalid_chars = set('\\/*?:[]')
    for wb in (internal_wb, client_wb):
        for name in wb.sheetnames:
            assert len(name) <= 31, f"{name!r} supera 31 caracteres"
            assert not (set(name) & invalid_chars), f"{name!r} tiene caracteres inválidos"


# --------------------------------------------------------------------------
# 9. Conteos consistentes con validation_report.yaml
# --------------------------------------------------------------------------

def test_conteos_de_pestanas_de_incidencia_coinciden_con_issues(fixture_run):
    ctx = load_run(fixture_run)
    wb = build_workbook(ctx, "internal")

    expected = {}
    for issue in FIXTURE_ISSUES:
        expected[issue["category"]] = expected.get(issue["category"], 0) + 1

    for code, count in expected.items():
        short_title = CATEGORY_CATALOG[code].short_title
        assert short_title in wb.sheetnames
        ws = wb[short_title]
        # La tabla empieza tras el bloque de explicación (7 filas + 1 en blanco);
        # contamos filas no vacías de la primera columna desde ahí en adelante.
        data_rows = 0
        found_header = False
        for row in ws.iter_rows():
            first_val = row[0].value
            if first_val == "ID histórico" or first_val == "Identificador histórico":
                found_header = True
                continue
            if found_header and first_val not in (None, ""):
                data_rows += 1
        assert data_rows == count, f"{code}: esperado {count}, encontrado {data_rows}"


# --------------------------------------------------------------------------
# 10. approved_for_enablon_import visible
# --------------------------------------------------------------------------

def test_approved_for_enablon_import_visible_en_ambas_versiones(internal_wb, client_wb):
    for wb in (internal_wb, client_wb):
        ws = wb["Resumen"]
        blob = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        assert "approved_for_enablon_import" in blob or "aprobación" in blob.lower() or "aprobado" in blob.lower()
        assert "false" in blob.lower() or "no implica aprobación" in blob.lower()


# --------------------------------------------------------------------------
# 11. No_Migra no contabilizado como error técnico
# --------------------------------------------------------------------------

def test_no_migra_tiene_severidad_no_tecnica(internal_wb):
    ws = internal_wb["No_Migra"]
    blob = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "No migra" in blob
    assert "error" not in blob.lower()


# --------------------------------------------------------------------------
# 12. StartingDate + Hora explicado
# --------------------------------------------------------------------------

def test_hallazgo_starting_date_hora_explicado(internal_wb, client_wb):
    for wb in (internal_wb, client_wb):
        ws = wb["Hallazgos"]
        blob = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        assert "Hora" in blob
        assert "StartingDate" in blob
        assert "Reference" in blob  # aclara explícitamente que no afecta a Reference.


# --------------------------------------------------------------------------
# 13. Fidelidad de entidad reflejada como pendiente
# --------------------------------------------------------------------------

def test_hallazgo_fidelidad_entidad_presente(internal_wb, client_wb):
    for wb in (internal_wb, client_wb):
        ws = wb["Hallazgos"]
        blob = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        assert "catálogo" in blob.lower()
        assert "confirmar" in blob.lower()


# --------------------------------------------------------------------------
# 14. Preguntas abiertas incluidas
# --------------------------------------------------------------------------

def test_preguntas_abiertas_incluidas(internal_wb):
    ws = internal_wb["Hallazgos"]
    blob = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    for qid in ("OQ-ETL-05", "OQ-ETL-06", "OQ-ENT-04", "OQ-ETL-02", "OQ-OBJ-07"):
        assert qid in blob


# --------------------------------------------------------------------------
# 15-17. Versión cliente sin hashes / paths / SQL
# --------------------------------------------------------------------------

def _all_text(wb) -> str:
    return " ".join(str(c.value) for sheet in wb.worksheets for row in sheet.iter_rows() for c in row if isinstance(c.value, str))


def test_cliente_sin_hashes(client_wb):
    blob = _all_text(client_wb)
    assert not re.search(r"\b[0-9a-f]{64}\b", blob)


def test_cliente_sin_paths_locales(client_wb):
    blob = _all_text(client_wb)
    assert not re.search(r"[A-Za-z]:\\Users", blob)
    assert "outputs/prototype" not in blob
    assert "outputs\\prototype" not in blob


def test_cliente_sin_sql(client_wb):
    blob = _all_text(client_wb)
    assert not re.search(r"\bSELECT\b", blob, re.IGNORECASE)
    assert "ITP_SIMULACRO" not in blob
    assert ".sql" not in blob


def test_cliente_sin_trazabilidad_tecnica(client_wb):
    assert "Trazabilidad" not in client_wb.sheetnames
    assert "Exportados" not in client_wb.sheetnames


# --------------------------------------------------------------------------
# 18. Versión interna con trazabilidad
# --------------------------------------------------------------------------

def test_interna_incluye_trazabilidad(internal_wb):
    assert "Trazabilidad" in internal_wb.sheetnames
    ws = internal_wb["Trazabilidad"]
    blob = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert re.search(r"\b[0-9a-f]{64}\b", blob)


# --------------------------------------------------------------------------
# 19. Comentario cliente presente en pestañas revisables
# --------------------------------------------------------------------------

def test_comentario_cliente_presente(client_wb):
    # "Dif_Historico" es una tabla agregada por columna (no una incidencia
    # por fila) -- no lleva "Comentario cliente", ver Fase 7.
    non_reviewable = {"Resumen", "Guía", "Hallazgos", CATEGORY_CATALOG["HISTORICAL_DIFFERENCES"].short_title}
    reviewable = [s for s in client_wb.sheetnames if s not in non_reviewable]
    assert reviewable, "El fixture debería producir al menos una pestaña de incidencia en la versión cliente."
    for name in reviewable:
        ws = client_wb[name]
        header_row = None
        for row in ws.iter_rows():
            if row[0].value in ("Identificador histórico",):
                header_row = [c.value for c in row]
                break
        assert header_row is not None, f"{name} no tiene cabecera de tabla reconocible"
        assert "Comentario cliente" in header_row


def test_decision_cliente_en_entidad_conflicto(client_wb):
    ws = client_wb["Entidad_Conflicto"]
    header_row = None
    for row in ws.iter_rows():
        if row[0].value == "Identificador histórico":
            header_row = [c.value for c in row]
            break
    assert header_row is not None
    assert "Decisión cliente" in header_row


# --------------------------------------------------------------------------
# 21-22. Workbook reabrible, sin fórmulas rotas
# --------------------------------------------------------------------------

def test_workbook_reabrible_sin_formulas(fixture_run, tmp_path):
    ctx = load_run(fixture_run)
    for audience in ("internal", "client"):
        wb = build_workbook(ctx, audience)
        path = save_workbook(wb, tmp_path / f"evidence_{audience}.xlsx")
        reopened = load_workbook(path)
        for sheet in reopened.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        assert not cell.value.startswith("="), f"Fórmula inesperada en {sheet.title}!{cell.coordinate}"


# --------------------------------------------------------------------------
# 23. Sin secretos (también en la versión interna -- ni contraseñas ni connection string)
# --------------------------------------------------------------------------

def test_interna_sin_contrasenas_ni_connection_string(internal_wb):
    blob = _all_text(internal_wb)
    assert not re.search(r"password", blob, re.IGNORECASE)
    assert not re.search(r"driver=", blob, re.IGNORECASE)


# --------------------------------------------------------------------------
# 24. CLI rechaza rutas fuera de outputs
# --------------------------------------------------------------------------

def test_cli_rechaza_ruta_fuera_de_outputs():
    runner = CliRunner()
    result = runner.invoke(cli, ["evidence", "drills", "--run", str(Path.cwd())])
    assert result.exit_code != 0
    assert "outputs" in (result.output or "") or "outputs" in str(result.exception or "")


def test_collector_rechaza_ejecucion_de_otro_objeto(tmp_path):
    run_dir = _build_fixture_run(tmp_path, run_id="fx-otro")
    vr_path = run_dir / "validation_report.yaml"
    data = yaml.safe_load(vr_path.read_text(encoding="utf-8"))
    data["run"]["migration_object"] = "Events"
    vr_path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")
    with pytest.raises(EvidenceSourceError):
        load_run(run_dir)


def test_collector_rechaza_manifiesto_sin_approved_flag(tmp_path):
    run_dir = _build_fixture_run(tmp_path)
    manifest_path = run_dir / "export_manifest.yaml"
    data = yaml.safe_load(manifest_path.read_text(encoding="utf-8"))
    del data["approved_for_enablon_import"]
    manifest_path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")
    with pytest.raises(EvidenceSourceError):
        load_run(run_dir)


def test_collector_rechaza_run_sin_issues_jsonl(tmp_path):
    run_dir = _build_fixture_run(tmp_path)
    (run_dir / "issues.jsonl").unlink()
    with pytest.raises(EvidenceSourceError):
        load_run(run_dir)


def test_collector_rechaza_carpeta_inexistente(tmp_path):
    with pytest.raises(EvidenceSourceError):
        resolve_run_dir(str(tmp_path / "no_existe"), runs_root=tmp_path)


# --------------------------------------------------------------------------
# 25. --audience both genera ambos archivos
# --------------------------------------------------------------------------

def test_cli_audience_both_genera_ambos(tmp_path):
    # `_generate_evidence` es la función que la CLI invoca tras resolver y
    # validar la ruta -- se prueba aquí directamente con una carpeta fuera
    # de outputs/ (permitido a este nivel) porque la restricción "dentro de
    # outputs/" ya se prueba por separado en
    # `test_cli_rechaza_ruta_fuera_de_outputs`, y crear el fixture dentro
    # del outputs/ real del repo ensuciaría una carpeta que no es de test.
    from src.cli import _generate_evidence

    run_dir = _build_fixture_run(tmp_path)
    written = _generate_evidence(run_dir, "both")
    assert {p.name for p in written} == {"evidence_internal.xlsx", "evidence_client.xlsx"}
    assert (run_dir / "evidence_internal.xlsx").is_file()
    assert (run_dir / "evidence_client.xlsx").is_file()


def test_cli_evidence_rechaza_audiencia_desconocida(tmp_path):
    run_dir = _build_fixture_run(tmp_path)
    runner = CliRunner()
    result = runner.invoke(cli, ["evidence", "drills", "--run", str(run_dir), "--audience", "foo"])
    assert result.exit_code != 0
