"""
Resolución correcta del destino Enablon en hojas de mapeo de campo (patrón de
8 columnas: `CampoOrigen | Field Destiny XML | Field Destiny ES | Adaptación |
Transformation From | (vacía) | XML | ES`).

Esta función NO sustituye ni modifica `read_field_mapping_sheet()` /
`FieldMapping` de `src/etl/excel_reader.py` -- es una implementación nueva e
independiente. La migración de `mapping_engine.py` hacia este resolver será
progresiva y en un incremento aparte (ver docstring de `excel_reader.py`).

Hallazgo confirmado con datos reales (MapeoBypass, MAP-AP, Mapeo_MOC_CT):
las columnas G ("XML") y H ("ES") de estas hojas NO están alineadas por fila
con las columnas A-E -- son una tabla de referencia independiente (el
catálogo completo de campos válidos de Enablon), colocada por casualidad en
el mismo rango de filas que la tabla de mapeo real. La columna B contiene
una fórmula `XLOOKUP(C, H:H, G:G)` que busca el texto de la columna C (de
esa fila) en TODA la columna H y devuelve el G correspondiente -- el destino
real de cada fila, no lo que hay en G de esa misma fila.

Este módulo reconstruye ese mismo comportamiento explícitamente:
- nunca lee la columna B (ni depende de si Excel recalculó la fórmula, ni de
  su valor cacheado);
- construye el diccionario ES -> XML a partir de TODAS las filas de las
  columnas G/H;
- detecta etiquetas ES repetidas con valores XML distintos (ambigüedad) y
  nunca elige una de forma silenciosa;
- normaliza solo espacio (strip + colapso de espacios repetidos), nunca
  mayúsculas/minúsculas;
- nunca inventa un destino cuando no hay correspondencia exacta.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

RESOLVED_EXACT = "resolved_exact"
UNRESOLVED = "unresolved"
AMBIGUOUS = "ambiguous"
MISSING_DESTINATION = "missing_destination"
INVALID_ROW = "invalid_row"

_ADAPTATION_TRUE = {"sí", "si", "yes", "true"}
_WHITESPACE = re.compile(r"\s+")


def normalize_label(text: str) -> str:
    """Normalización de espacio "segura": strip + colapso de espacios
    internos repetidos. Deliberadamente NO toca mayúsculas/minúsculas -- una
    etiqueta que solo difiere en case podría ser un error real de captura,
    no algo que deba fusionarse sin más."""
    return _WHITESPACE.sub(" ", text.strip())


@dataclass
class EnablonReferenceCatalog:
    """Catálogo ES -> XML construido a partir de columnas G/H de una hoja.
    `raw_entries` conserva el orden de aparición para trazabilidad; nunca se
    usa como fuente de verdad para la resolución (esa es `es_to_xml`, ya
    filtrada de ambigüedades)."""
    es_to_xml: dict[str, str] = field(default_factory=dict)
    duplicate_es_labels: dict[str, list[str]] = field(default_factory=dict)
    raw_entries: list[tuple[str, str]] = field(default_factory=list)  # (xml, es)


def build_reference_catalog(rows: list[tuple]) -> EnablonReferenceCatalog:
    """Construye el catálogo a partir de TODAS las filas de una hoja
    (columnas índice 6 = XML, 7 = ES) -- nunca de la columna B (fórmula)."""
    catalog = EnablonReferenceCatalog()
    seen: dict[str, set[str]] = {}

    for row in rows:
        xml_raw = row[6] if len(row) > 6 else None
        es_raw = row[7] if len(row) > 7 else None
        if xml_raw is None or es_raw is None:
            continue
        xml_norm = normalize_label(str(xml_raw))
        es_norm = normalize_label(str(es_raw))
        if not xml_norm or not es_norm:
            continue
        catalog.raw_entries.append((xml_norm, es_norm))
        seen.setdefault(es_norm, set()).add(xml_norm)

    for es_norm, xml_values in seen.items():
        if len(xml_values) == 1:
            catalog.es_to_xml[es_norm] = next(iter(xml_values))
        else:
            catalog.duplicate_es_labels[es_norm] = sorted(xml_values)

    return catalog


@dataclass
class ResolvedFieldMapping:
    source_field: str
    source_table: str | None
    destination_label_es: str | None
    destination_field_xml: str | None
    transformation_rule: str | None
    adaptation: bool
    resolution_status: str
    resolution_method: str
    source_file: str
    source_sheet: str
    source_row: int
    warnings: list[str] = field(default_factory=list)


def _row_has_mapping_content(row: tuple) -> bool:
    """Comprueba contenido específicamente en las columnas A-E (0-4), el
    rango de la tabla de mapeo real. Una fila puede tener contenido SOLO en
    G/H (6,7) -- es una fila del catálogo de referencia que se quedó sin
    pareja en la tabla de mapeo (el catálogo suele tener más entradas que la
    tabla de mapeo, ver MAP-AP: 73 entradas de catálogo frente a 36 filas de
    mapeo) -- eso no es una fila de mapeo inválida, no se reporta como tal."""
    return any(v is not None and str(v).strip() != "" for v in row[:5])


def _resolve_row(
    row: tuple,
    row_index: int,
    catalog: EnablonReferenceCatalog,
    source_file: str,
    source_sheet: str,
) -> ResolvedFieldMapping | None:
    if not _row_has_mapping_content(row):
        return None  # fila vacía en A-E: o está totalmente vacía, o es una
        # fila "solo catálogo" (G/H) -- ninguna de las dos es una fila de
        # mapeo a reportar.

    campo_origen_raw = row[0] if len(row) > 0 else None
    if campo_origen_raw is None or str(campo_origen_raw).strip() == "":
        return ResolvedFieldMapping(
            source_field="",
            source_table=None,
            destination_label_es=None,
            destination_field_xml=None,
            transformation_rule=None,
            adaptation=False,
            resolution_status=INVALID_ROW,
            resolution_method="no_source_field",
            source_file=source_file,
            source_sheet=source_sheet,
            source_row=row_index,
            warnings=["Fila con contenido pero sin CampoOrigen -- posible fila corrupta o de formato."],
        )

    campo_origen = str(campo_origen_raw)
    adaptacion_raw = str(row[3]).strip().lower() if len(row) > 3 and row[3] is not None else "no"
    adaptation = adaptacion_raw in _ADAPTATION_TRUE
    regla_raw = row[4] if len(row) > 4 else None
    transformation_rule = str(regla_raw) if regla_raw is not None and str(regla_raw).strip() else None

    destino_es_raw = row[2] if len(row) > 2 else None
    if destino_es_raw is None or str(destino_es_raw).strip() == "":
        return ResolvedFieldMapping(
            source_field=campo_origen,
            source_table=None,
            destination_label_es=None,
            destination_field_xml=None,
            transformation_rule=transformation_rule,
            adaptation=adaptation,
            resolution_status=MISSING_DESTINATION,
            resolution_method="no_destination_label",
            source_file=source_file,
            source_sheet=source_sheet,
            source_row=row_index,
        )

    destino_es = normalize_label(str(destino_es_raw))

    if destino_es in catalog.duplicate_es_labels:
        candidates = catalog.duplicate_es_labels[destino_es]
        return ResolvedFieldMapping(
            source_field=campo_origen,
            source_table=None,
            destination_label_es=destino_es,
            destination_field_xml=None,
            transformation_rule=transformation_rule,
            adaptation=adaptation,
            resolution_status=AMBIGUOUS,
            resolution_method="es_label_duplicated_in_catalog",
            source_file=source_file,
            source_sheet=source_sheet,
            source_row=row_index,
            warnings=[
                f"La etiqueta ES {destino_es!r} aparece en el catálogo con "
                f"{len(candidates)} valores XML distintos: {candidates} -- "
                "no se elige ninguno automáticamente."
            ],
        )

    if destino_es in catalog.es_to_xml:
        return ResolvedFieldMapping(
            source_field=campo_origen,
            source_table=None,
            destination_label_es=destino_es,
            destination_field_xml=catalog.es_to_xml[destino_es],
            transformation_rule=transformation_rule,
            adaptation=adaptation,
            resolution_status=RESOLVED_EXACT,
            resolution_method="exact_lookup_es_to_xml",
            source_file=source_file,
            source_sheet=source_sheet,
            source_row=row_index,
        )

    return ResolvedFieldMapping(
        source_field=campo_origen,
        source_table=None,
        destination_label_es=destino_es,
        destination_field_xml=None,
        transformation_rule=transformation_rule,
        adaptation=adaptation,
        resolution_status=UNRESOLVED,
        resolution_method="no_match_in_catalog",
        source_file=source_file,
        source_sheet=source_sheet,
        source_row=row_index,
        warnings=[
            f"La etiqueta ES {destino_es!r} no aparece en el catálogo de "
            "referencia (columnas XML/ES) de esta hoja."
        ],
    )


def resolve_field_mapping_sheet(path: str | Path, sheet_name: str) -> list[ResolvedFieldMapping]:
    """Lee una hoja del patrón de mapeo de campo y resuelve el destino XML
    real de cada fila contra el catálogo embebido en la propia hoja.

    `source_table` queda siempre en `None`: esta función no tiene forma
    fiable de saber la tabla SQL de origen de un campo -- eso requiere
    cruzar contra el análisis de columnas SELECT de las queries (fuera del
    alcance de este resolver), y no se inventa aquí.
    """
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        return []

    data_rows = rows[1:]  # fila 0 = cabecera
    catalog = build_reference_catalog(data_rows)

    resolved: list[ResolvedFieldMapping] = []
    for i, row in enumerate(data_rows, start=1):
        item = _resolve_row(row, i, catalog, source_file=str(path), source_sheet=sheet_name)
        if item is not None:
            resolved.append(item)
    return resolved
